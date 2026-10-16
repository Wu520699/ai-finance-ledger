#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AI 票据智能台账 v0.1
场景：小微餐饮门店（烤馒头店）业财融合实践
输入：data/raw_tickets.csv （AI 从票据 / 收款截图识别出的结构化数据）
输出：out/台账.xlsx（明细台账 / 科目汇总 / 日汇总 / 经营小结）
"""
import argparse
import csv
import math
import os
import sys
from collections import defaultdict
from datetime import datetime

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, "data", "raw_tickets.csv")
OUT_DIR = os.path.join(BASE, "out")
OUT_XLSX = os.path.join(OUT_DIR, "台账.xlsx")
OUT_TXT = os.path.join(OUT_DIR, "经营小结.txt")
DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日")
MONEY_FMT = "#,##0.00"
PERCENT_FMT = "0.0%"
os.makedirs(OUT_DIR, exist_ok=True)

# ---------- 科目映射规则：小微餐饮的最小科目集 ----------
EXPENSE_RULES = [
    (("面粉", "酵母", "芝麻", "豆沙", "馅", "糖", "食用油", "调和油", "牛奶", "黄油", "原料", "食材"), "主营业务成本-食材"),
    (("包装", "纸袋", "餐盒", "一次性", "袋子"), "主营业务成本-包装耗材"),
    (("电费", "水费", "燃气", "煤气", "充电"), "管理费用-水电燃气"),
    (("房租", "租金", "摊位", "物业"), "销售费用-房租"),
    (("工资", "临时工", "兼职", "人工"), "销售费用-人工"),
    (("美团", "饿了么", "佣金", "平台服务费"), "销售费用-平台佣金"),
    (("运费", "快递", "配送"), "销售费用-配送费"),
    (("设备", "维修", "烤炉", "周转箱"), "管理费用-设备维修"),
    (("税", "发票"), "税金及附加"),
]
OTHER_EXPENSE = "其他支出"
INCOME_ACCOUNT = "主营业务收入"


def classify(direction, summary):
    if direction == "收":
        return INCOME_ACCOUNT
    for keys, acc in EXPENSE_RULES:
        for k in keys:
            if k in summary:
                return acc
    return OTHER_EXPENSE


def parse_amount(text):
    """宽口径解析金额：去空白 / 千分位 / 常见币种符号；非法或非有限值返回 None。"""
    s = str(text if text is not None else "").replace(",", "").replace("，", "").strip()
    for ch in ("￥", "¥", "元", "$"):
        s = s.replace(ch, "")
    s = s.strip()
    if not s:
        return None
    try:
        value = float(s)
    except ValueError:
        return None
    if not math.isfinite(value):
        return None
    return value


def normalize_date(text):
    """把常见日期写法统一成 YYYY-MM-DD；无法识别时原样返回，交由调用方处理。"""
    s = (text or "").strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return s


def load(src=SRC):
    rows = []
    skipped = []
    with open(src, "r", encoding="utf-8-sig", newline="") as f:
        for idx, r in enumerate(csv.DictReader(f), start=2):
            r = r or {}
            raw_date = (r.get("日期") or "").strip()
            summary = (r.get("摘要") or "").strip()
            direction = (r.get("收付") or "").strip()
            if not raw_date:
                continue
            amount = parse_amount(r.get("金额"))
            if amount is None:
                skipped.append(idx)
                continue
            rows.append({
                "日期": normalize_date(raw_date),
                "摘要": summary,
                "收付": direction,
                "金额": amount,
                "支付方式": (r.get("支付方式") or "").strip(),
                "来源图片": (r.get("来源图片") or "").strip(),
                "科目": classify(direction, summary),
            })
    if skipped:
        sys.stderr.write("警告：第 %s 行金额缺失或无法解析，已跳过。\n"
                         % "、".join(str(i) for i in skipped))
    rows.sort(key=lambda x: x["日期"])
    return rows


def summarize(rows):
    income = sum(r["金额"] for r in rows if r["收付"] == "收")
    expense = sum(r["金额"] for r in rows if r["收付"] == "支")
    by_acc = defaultdict(float)
    for r in rows:
        if r["收付"] == "支":
            by_acc[r["科目"]] += r["金额"]
    by_day = defaultdict(lambda: [0.0, 0.0])
    for r in rows:
        if r["收付"] == "收":
            by_day[r["日期"]][0] += r["金额"]
        else:
            by_day[r["日期"]][1] += r["金额"]
    food = by_acc.get("主营业务成本-食材", 0.0)
    gross = income - food
    return {
        "income": income, "expense": expense, "profit": income - expense,
        "gross": gross, "food": food, "by_acc": dict(by_acc), "by_day": dict(by_day),
        "days": len(by_day),
    }


def write_excel(rows, s, path=OUT_XLSX):
    out_dir = os.path.dirname(os.path.abspath(path))
    os.makedirs(out_dir, exist_ok=True)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        csv_path = os.path.join(out_dir, "台账_明细.csv")
        fields = ["日期", "摘要", "收付", "金额", "支付方式", "来源图片", "科目"]
        with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            w.writerows(rows)
        return None
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")

    def style(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for c in ws[1]:
            c.font = head_font
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center")

    def numfmt(ws, cols, fmt):
        """给指定列（字母）去掉表头后的单元格设置数字格式。"""
        for col in cols:
            for cell in list(ws[col])[1:]:
                cell.number_format = fmt

    ws = wb.active
    ws.title = "明细台账"
    ws.append(["日期", "摘要", "收/支", "金额(元)", "会计科目", "支付方式", "AI识别源文件"])
    for r in rows:
        ws.append([r["日期"], r["摘要"], r["收付"], r["金额"], r["科目"], r["支付方式"], r["来源图片"]])
    style(ws, [12, 26, 8, 12, 24, 12, 20])
    numfmt(ws, ["D"], MONEY_FMT)

    ws2 = wb.create_sheet("科目汇总")
    ws2.append(["会计科目", "金额(元)", "占支出比"])
    for acc, amt in sorted(s["by_acc"].items(), key=lambda x: -x[1]):
        ws2.append([acc, round(amt, 2), (amt / s["expense"]) if s["expense"] else 0])
    ws2.append(["支出合计", round(s["expense"], 2), 1.0 if s["expense"] else 0])
    style(ws2, [26, 14, 12])
    numfmt(ws2, ["B"], MONEY_FMT)
    numfmt(ws2, ["C"], PERCENT_FMT)

    ws3 = wb.create_sheet("日汇总")
    ws3.append(["日期", "收入(元)", "支出(元)", "当日净额(元)"])
    for d in sorted(s["by_day"]):
        inc, exp = s["by_day"][d]
        ws3.append([d, round(inc, 2), round(exp, 2), round(inc - exp, 2)])
    style(ws3, [14, 14, 14, 16])
    numfmt(ws3, ["B", "C", "D"], MONEY_FMT)

    ws4 = wb.create_sheet("经营小结")
    for line in report_lines(s):
        ws4.append([line])
    ws4.column_dimensions["A"].width = 60

    wb.save(path)
    return path


def report_lines(s):
    days = max(s["days"], 1)
    lines = [
        "小微餐饮门店经营小结（AI 票据智能台账自动生成）",
        "=" * 46,
        "统计天数：%d 天" % s["days"],
        "营业额合计：%.2f 元（日均 %.2f 元）" % (s["income"], s["income"] / days),
        "支出合计：%.2f 元（日均 %.2f 元）" % (s["expense"], s["expense"] / days),
        "净现金流：%.2f 元" % s["profit"],
        "食材成本：%.2f 元，占营业额 %.1f%%" % (s["food"], s["food"] / s["income"] * 100 if s["income"] else 0),
        "毛利（营业额-食材）：%.2f 元，毛利率 %.1f%%" % (s["gross"], s["gross"] / s["income"] * 100 if s["income"] else 0),
        "-" * 46,
        "成本费用结构：",
    ]
    for acc, amt in sorted(s["by_acc"].items(), key=lambda x: -x[1]):
        lines.append("  %-22s %8.2f 元  %5.1f%%" % (acc, amt, amt / s["expense"] * 100 if s["expense"] else 0))
    lines += [
        "-" * 46,
        "业财提示：",
        "  1. 毛利率低于 60% 时，优先查面粉等主料单价与出成率，而不是直接涨价。",
        "  2. 平台佣金与配送费单独归集，单量越高越要盯，它是隐形毛利杀手。",
        "  3. 现金收款比例高的门店，建议每日对一次收款码流水，防止账实不符。",
        "（本小结由脚本按日流水自动生成，可作为门店经营复盘的底稿）",
    ]
    return lines


def main():
    ap = argparse.ArgumentParser(
        description="AI 票据智能台账 v0.1：读取票据 CSV，按科目自动入账，输出 Excel 台账与经营小结。",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    ap.add_argument("--src", default=SRC, help="票据流水 CSV 路径")
    ap.add_argument("--out", default=OUT_XLSX, help="台账 Excel 输出路径")
    args = ap.parse_args()

    rows = load(args.src)
    s = summarize(rows)
    path = write_excel(rows, s, args.out)
    os.makedirs(os.path.dirname(os.path.abspath(OUT_TXT)), exist_ok=True)
    with open(OUT_TXT, "w", encoding="utf-8") as f:
        f.write("\n".join(report_lines(s)))
    print("票据条数：", len(rows))
    print("统计天数：", s["days"])
    print("营业额：%.2f  支出：%.2f  净额：%.2f" % (s["income"], s["expense"], s["profit"]))
    print("毛利率：%.1f%%" % (s["gross"] / s["income"] * 100 if s["income"] else 0))
    print("Excel：", path)
    print("小结：", OUT_TXT)


if __name__ == "__main__":
    main()
